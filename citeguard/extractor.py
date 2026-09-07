"""Document reading, paragraph splitting, and citation detection."""

from __future__ import annotations

import re
from collections.abc import Iterator
from pathlib import Path
from typing import TYPE_CHECKING

from docx import Document
from docx.document import Document as DocxDocument
from docx.oxml.ns import qn
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph

from .models import ExistingCitation, ParsedDocument, Sentence
from .retrieval import normalize_doi
from .sentence_splitter import split_sentences

if TYPE_CHECKING:
    from .models import EnrichedDocument

DATE_TEXT_PATTERN = r"(?:\d{4}[a-zA-Z]?|(?i:t\.?\s*y\.?|n\.?\s*d\.?))"
AUTHOR_WORD_PATTERN = r"[A-ZÇĞİÖŞÜ][A-Za-zÀ-ÖØ-öø-ÿÇĞİÖŞÜçğıöşü'’.-]*"
PARENTHETICAL_RE = re.compile(r"\((?P<body>[^()\n]{1,250})\)")
PARENTHETICAL_PART_RE = re.compile(
    rf"^(?P<authors>.+?)\s*,\s*(?P<date>{DATE_TEXT_PATTERN})"
    r"(?:\s*,\s*(?:p{1,2}|s{1,2})\.\s*\d+(?:\s*[-–]\s*\d+)?)?\s*$",
    re.IGNORECASE,
)
NARRATIVE_AUTHOR_YEAR_RE = re.compile(
    rf"\b(?P<authors>{AUTHOR_WORD_PATTERN}"
    rf"(?:(?:,?\s+)(?:et\s+al\.|vd\.|ve|and|&|{AUTHOR_WORD_PATTERN})){{0,10}})"
    rf"\s+\((?P<date>{DATE_TEXT_PATTERN})"
    r"(?:\s*,\s*(?:[pP]{1,2}|[sS]{1,2})\.\s*\d+(?:\s*[-–]\s*\d+)?)?\)",
)
NUMBERED_RE = re.compile(r"\[(?P<numbers>\d+(?:\s*[-–]\s*\d+|\s*,\s*\d+)*)\]")
DOI_RE = re.compile(r"\b10\.\d{4,9}/[-._;()/:A-Z0-9]+", re.IGNORECASE)

SUPPORTED_SUFFIXES = {".md", ".txt", ".docx"}


def read_paragraphs(path: str | Path) -> list[str]:
    file_path = Path(path)
    suffix = file_path.suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        raise ValueError(
            f"Unsupported file type: {suffix or '<none>'}. "
            f"Supported types: {', '.join(sorted(SUPPORTED_SUFFIXES))}."
        )

    if suffix == ".docx":
        doc = Document(file_path)
        return _read_docx_paragraphs(doc)

    text = file_path.read_text(encoding="utf-8")
    return split_text_paragraphs(text)


def _read_docx_paragraphs(document: DocxDocument) -> list[str]:
    """Read body and table-cell paragraphs in their document order."""
    return [text for block in _iter_docx_blocks(document) if (text := block.text.strip())]


def _iter_docx_blocks(parent: DocxDocument | _Cell) -> Iterator[Paragraph]:
    """Yield paragraphs recursively, without duplicating merged table cells."""
    parent_element = parent.element.body if isinstance(parent, DocxDocument) else parent._tc
    for child in parent_element.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, parent)
        elif child.tag == qn("w:tbl"):
            table = Table(child, parent)
            seen_cells: set[object] = set()
            for row in table.rows:
                for cell in row.cells:
                    cell_identity = cell._tc
                    if cell_identity in seen_cells:
                        continue
                    seen_cells.add(cell_identity)
                    yield from _iter_docx_blocks(cell)


def split_text_paragraphs(text: str) -> list[str]:
    normalized = text.replace("\r\n", "\n").replace("\r", "\n")
    blocks = re.split(r"\n\s*\n", normalized)
    paragraphs: list[str] = []
    for block in blocks:
        block = block.strip()
        if not block:
            continue
        # Preserve Markdown headings as their own logical paragraph when possible.
        lines = [line.strip() for line in block.splitlines() if line.strip()]
        if len(lines) > 1 and lines[0].startswith("#"):
            paragraphs.append(lines[0])
            rest = " ".join(lines[1:]).strip()
            if rest:
                paragraphs.append(rest)
        else:
            paragraphs.append(" ".join(lines))
    return paragraphs


def extract_citations(paragraphs: list[str]) -> list[ExistingCitation]:
    citations: list[ExistingCitation] = []

    for paragraph_index, paragraph in enumerate(paragraphs):
        occupied: list[tuple[int, int]] = []

        for match in PARENTHETICAL_RE.finditer(paragraph):
            matched_parts = []
            for part in match.group("body").split(";"):
                part_match = PARENTHETICAL_PART_RE.fullmatch(part.strip())
                if part_match and _looks_like_author(part_match.group("authors")):
                    matched_parts.append(part_match)
            if not matched_parts:
                continue
            for part_match in matched_parts:
                date_text = part_match.group("date")
                no_date = _is_no_date(date_text)
                citations.append(
                    ExistingCitation(
                        raw_text=match.group(0),
                        authors=part_match.group("authors").strip(),
                        year=None if no_date else int(date_text[:4]),
                        doi=None,
                        numbered_ref=None,
                        paragraph_index=paragraph_index,
                        char_offset=match.start(),
                        no_date=no_date,
                    )
                )
            occupied.append(match.span())

        for match in NARRATIVE_AUTHOR_YEAR_RE.finditer(paragraph):
            if any(start <= match.start() < end for start, end in occupied):
                continue
            date_text = match.group("date")
            no_date = _is_no_date(date_text)
            citations.append(
                ExistingCitation(
                    raw_text=match.group(0),
                    authors=match.group("authors").strip(),
                    year=None if no_date else int(date_text[:4]),
                    doi=None,
                    numbered_ref=None,
                    paragraph_index=paragraph_index,
                    char_offset=match.start(),
                    no_date=no_date,
                )
            )
            occupied.append(match.span())

        for match in NUMBERED_RE.finditer(paragraph):
            numbers = _expand_numbered_citations(match.group("numbers"))
            for number in numbers:
                citations.append(
                    ExistingCitation(
                        raw_text=match.group(0),
                        authors=None,
                        year=None,
                        doi=None,
                        numbered_ref=number,
                        paragraph_index=paragraph_index,
                        char_offset=match.start(),
                    )
                )
            occupied.append(match.span())

        for match in DOI_RE.finditer(paragraph):
            if any(start <= match.start() < end for start, end in occupied):
                continue
            citations.append(
                ExistingCitation(
                    raw_text=match.group(0),
                    authors=None,
                    year=None,
                    doi=normalize_doi(match.group(0)),
                    numbered_ref=None,
                    paragraph_index=paragraph_index,
                    char_offset=match.start(),
                )
            )

    return sorted(citations, key=lambda c: (c.paragraph_index, c.char_offset))


def _looks_like_author(value: str) -> bool:
    author = value.strip()
    if not author or not author[0].isupper() or len(author.split()) > 20:
        return False
    return bool(re.fullmatch(r"[\wÀ-ÖØ-öø-ÿÇĞİÖŞÜçğıöşü .,'’&-]+", author))


def _is_no_date(value: str) -> bool:
    return not value.strip()[0].isdigit()


_RANGE_SEP_RE = re.compile(r"\s*[-–]\s*")
_MAX_RANGE_SIZE = 100


def _expand_numbered_citations(text: str) -> list[int]:
    numbers: list[int] = []
    for part in text.split(","):
        part = part.strip()
        if not part:
            continue
        range_match = _RANGE_SEP_RE.split(part, maxsplit=1)
        if len(range_match) == 2:
            try:
                start = int(range_match[0].strip())
                end = int(range_match[1].strip())
            except ValueError:
                continue
            if start > end:
                start, end = end, start
            if end - start > _MAX_RANGE_SIZE:
                end = start + _MAX_RANGE_SIZE
            numbers.extend(range(start, end + 1))
        else:
            try:
                numbers.append(int(part))
            except ValueError:
                continue
    seen: set[int] = set()
    unique: list[int] = []
    for n in numbers:
        if n not in seen and n > 0:
            seen.add(n)
            unique.append(n)
    return sorted(unique)


def parse_enriched_document(path: str | Path) -> EnrichedDocument:
    from .bibliography import parse_bibliography
    from .models import EnrichedDocument as _EnrichedDocument

    paragraphs = read_paragraphs(path)
    bibliography_start, entries = parse_bibliography(paragraphs)
    content_paragraphs = (
        paragraphs[:bibliography_start] if bibliography_start is not None else paragraphs
    )
    citations = extract_citations(content_paragraphs)

    # Build enriched structure with sentences
    all_sentences: list[Sentence] = []
    all_paragraphs: list = []
    bib_start_idx = bibliography_start

    for pidx, para_text in enumerate(paragraphs):
        is_bib = bib_start_idx is not None and pidx >= bib_start_idx
        sentences = split_sentences(
            para_text,
            paragraph_index=pidx,
            is_bibliography=is_bib,
        )
        for _sidx, sent in enumerate(sentences):
            # Link citations that fall within this sentence's offset range
            sent_citations = [
                c for c in citations
                if c.paragraph_index == pidx
                and sent.start_offset <= c.char_offset < sent.end_offset
            ]
            sent.citations = sent_citations
            all_sentences.append(sent)
        all_paragraphs.append(type("Paragraph", (), {
            "text": para_text,
            "index": pidx,
            "sentences": sentences,
            "is_bibliography": is_bib,
        })())

    return _EnrichedDocument(
        path=str(Path(path)),
        paragraphs=all_paragraphs,
        sentences=all_sentences,
        citations=citations,
        bibliography_entries=entries,
        bibliography_start_index=bib_start_idx,
    )


def parse_document(path: str | Path) -> ParsedDocument:
    from .bibliography import parse_bibliography

    paragraphs = read_paragraphs(path)
    bibliography_start, entries = parse_bibliography(paragraphs)
    content_paragraphs = (
        paragraphs[:bibliography_start] if bibliography_start is not None else paragraphs
    )
    citations = extract_citations(content_paragraphs)
    return ParsedDocument(
        path=str(Path(path)),
        paragraphs=paragraphs,
        citations=citations,
        bibliography_entries=entries,
        bibliography_start_index=bibliography_start,
    )
