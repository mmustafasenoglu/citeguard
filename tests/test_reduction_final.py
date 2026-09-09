"""End-to-end Step 8 attribution-reduction safety tests."""

from __future__ import annotations

from pathlib import Path
from unittest.mock import MagicMock

import pytest
from docx import Document

from citeguard.audit import (
    AuditOptions,
    AuditResult,
    CitedSourceAssessment,
    ClaimAssessment,
    ExecutionContext,
)
from citeguard.models import (
    Claim,
    ClaimType,
    Evidence,
    EvidenceType,
    ExistingCitation,
    Severity,
    SourceCandidate,
    Verdict,
    VerificationStatus,
)
from citeguard.reduction import (
    MeaningVerdict,
    ReductionOptions,
    RewriteCandidate,
    improve_attribution,
    rank_candidates,
)
from citeguard.reduction.apply import TextReplacement
from citeguard.reduction.docx_apply import (
    UnsafeDocxStructureError,
    write_revised_docx,
)
from citeguard.reduction.grounding import build_grounded_reduction_requests
from citeguard.reduction.meaning import EntailmentDirection
from citeguard.reduction.metrics import compute_reduction_metrics
from citeguard.reduction.report import reduction_result_report
from citeguard.reduction.source_evaluator import evaluate_source_overlap
from citeguard.rewrite import RewriteResult, RewriteStatus
from citeguard.similarity.models import SimilarityEngineResult

ORIGINAL = "The intervention significantly reduced mortality by 12% (Smith, 2024)."
REWRITE = "Following the intervention, mortality fell by 12% (Smith, 2024)."
SOURCE = "The intervention significantly reduced mortality by 12%."


class _PreservedSemantic:
    def similarity(self, _original: str, _candidate: str) -> float:
        return 0.96


class _PreservedEntailment:
    def evaluate(self, _premise: str, _hypothesis: str) -> EntailmentDirection:
        return EntailmentDirection(0.95, MeaningVerdict.PRESERVED)


class _RewriteProvider:
    def __init__(self, text: str = REWRITE) -> None:
        self.text = text
        self.calls = 0

    def rewrite(self, request) -> RewriteResult:
        self.calls += 1
        return RewriteResult(
            original_text=request.context.original_text,
            rewritten_text=self.text,
            status=RewriteStatus.SUCCESS,
            provider="fake",
            model="fixed",
        )


def _audit(path: Path, _options) -> AuditResult:
    citation = ExistingCitation(
        "(Smith, 2024)", "Smith", 2024, None, None, 0, ORIGINAL.index("(")
    )
    claim = Claim(
        text=ORIGINAL,
        search_query=ORIGINAL,
        claim_type=ClaimType.STATISTIC,
        severity=Severity.HIGH,
        paragraph_index=0,
        has_existing_citation=True,
        linked_citation=citation,
    )
    candidate = SourceCandidate(
        title="Trial",
        authors=["Smith"],
        year=2024,
        venue="Journal",
        doi="10.1000/trial",
        url=None,
        abstract=SOURCE,
        source_api="fake",
    )
    evidence = Evidence(
        text=SOURCE,
        source_title="Trial",
        source_api="fake",
        evidence_type=EvidenceType.ABSTRACT,
        verdict=Verdict.SUPPORTED,
    )
    assessment = ClaimAssessment(
        claim,
        [
            CitedSourceAssessment(
                citation_raw=citation.raw_text,
                entry_index=0,
                bib_status=VerificationStatus.VERIFIED,
                candidate=candidate,
                evaluated=True,
                reason="fixture",
                verdict=Verdict.SUPPORTED,
                evidence=[evidence],
            )
        ],
        "supported",
    )
    return AuditResult(
        document=str(path),
        claims=[claim],
        bibliography_verification=[],
        claim_assessments=[assessment],
        suggestions=[],
        bibliography_issues=[],
        similarity=None,
        metrics=MagicMock(),
        review_queue=[],
        diagnostics=[],
        execution=ExecutionContext(False, True, False),
    )


def _write_fixture(tmp_path: Path, suffix: str = ".md") -> tuple[Path, Path]:
    paper = tmp_path / f"paper{suffix}"
    corpus = tmp_path / "corpus.txt"
    paper.write_text(
        f"{ORIGINAL}\n\nReferences\n\nSmith, J. (2024). Trial.", encoding="utf-8"
    )
    corpus.write_text(SOURCE, encoding="utf-8")
    return paper, corpus


def test_source_aware_overlap_measures_actual_source_reduction() -> None:
    candidate = RewriteCandidate(REWRITE, "fake")
    evaluate_source_overlap(ORIGINAL, candidate, SOURCE)
    assert candidate.source_overlap_improved is True
    assert candidate.source_exact_overlap_after < candidate.source_exact_overlap_before
    assert (
        candidate.source_lexical_similarity_after
        < candidate.source_lexical_similarity_before
    )


def test_original_difference_does_not_substitute_for_source_reduction() -> None:
    candidate = RewriteCandidate(SOURCE, "fake")
    evaluate_source_overlap("Completely different original wording here.", candidate, SOURCE)
    assert candidate.source_overlap_improved is False
    assert candidate.rejection_reasons


def test_unvalidated_or_source_unimproved_candidate_cannot_rank() -> None:
    candidate = RewriteCandidate(
        REWRITE,
        "fake",
        citations_preserved=True,
        numeric_integrity=True,
        factual_integrity=True,
        meaning_verdict=MeaningVerdict.PRESERVED,
    )
    assert rank_candidates([candidate]) == []


def test_end_to_end_markdown_apply_and_actual_rescan(tmp_path: Path) -> None:
    paper, corpus = _write_fixture(tmp_path)
    output = tmp_path / "paper.revised.md"
    original_bytes = paper.read_bytes()
    provider = _RewriteProvider()
    result = improve_attribution(
        paper,
        ReductionOptions(
            corpus=corpus,
            corpus_license="CC0",
            apply=True,
            output=output,
            candidate_count=1,
            max_iterations=2,
            rewrite_provider=provider,
            semantic_backend=_PreservedSemantic(),
            entailment_backend=_PreservedEntailment(),
            audit_runner=_audit,
        ),
    )
    assert paper.read_bytes() == original_bytes
    assert output.exists()
    assert REWRITE in output.read_text(encoding="utf-8")
    assert "Smith, J. (2024). Trial." in output.read_text(encoding="utf-8")
    assert result.after.overall_similarity_pct < result.before.overall_similarity_pct
    assert result.metrics.rewritten_passages == 1
    assert provider.calls == 1
    report = reduction_result_report(result)
    assert report["schema_version"] == "2"
    assert report["privacy"]["source_overwritten"] is False


def test_preview_is_non_destructive_and_does_not_invoke_rewriter(tmp_path: Path) -> None:
    paper, corpus = _write_fixture(tmp_path)
    provider = MagicMock()
    before = paper.read_bytes()
    result = improve_attribution(
        paper,
        ReductionOptions(corpus=corpus, corpus_license="CC0", rewrite_provider=provider),
    )
    assert result.stop_reason == "preview_only"
    assert paper.read_bytes() == before
    provider.rewrite.assert_not_called()


def test_apply_requires_distinct_output(tmp_path: Path) -> None:
    paper, corpus = _write_fixture(tmp_path)
    with pytest.raises(ValueError, match="output path"):
        improve_attribution(paper, ReductionOptions(corpus=corpus, apply=True))
    with pytest.raises(ValueError, match="differ"):
        improve_attribution(
            paper, ReductionOptions(corpus=corpus, apply=True, output=paper)
        )


def test_cli_apply_requires_output_and_conflicts_with_dry_run(tmp_path: Path) -> None:
    from click.testing import CliRunner

    from citeguard.cli import main

    paper, corpus = _write_fixture(tmp_path)
    base = [
        "improve-attribution",
        str(paper),
        "--corpus",
        str(corpus),
        "--corpus-license",
        "CC0",
    ]
    missing = CliRunner().invoke(main, [*base, "--apply"])
    conflict = CliRunner().invoke(
        main, [*base, "--apply", "--dry-run", "--output", str(tmp_path / "out.md")]
    )
    assert missing.exit_code == 2
    assert "requires --output" in missing.output
    assert conflict.exit_code == 2
    assert "cannot be used together" in conflict.output


def test_offline_credentials_never_call_rewriter(tmp_path: Path, monkeypatch) -> None:
    paper, corpus = _write_fixture(tmp_path)
    output = tmp_path / "offline.md"
    monkeypatch.setenv("ANTHROPIC_API_KEY", "fake-credential")
    provider = MagicMock()
    result = improve_attribution(
        paper,
        ReductionOptions(
            corpus=corpus,
            corpus_license="CC0",
            apply=True,
            output=output,
            offline=True,
            rewrite_provider=provider,
            audit_runner=_audit,
        ),
    )
    provider.rewrite.assert_not_called()
    assert result.stop_reason == "no_safe_candidates"
    assert output.read_bytes() == paper.read_bytes()


def test_unsupported_added_claim_is_rejected_by_document_pipeline(tmp_path: Path) -> None:
    paper, corpus = _write_fixture(tmp_path)
    output = tmp_path / "unsupported.md"

    class SelectiveEntailment(_PreservedEntailment):
        def evaluate(self, premise: str, hypothesis: str) -> EntailmentDirection:
            if "cost" in hypothesis.casefold() and "cost" not in premise.casefold():
                return EntailmentDirection(0.2, MeaningVerdict.UNKNOWN)
            return super().evaluate(premise, hypothesis)

    provider = _RewriteProvider(
        "Following the intervention, mortality fell by 12% (Smith, 2024). "
        "The treatment also reduced cost."
    )
    result = improve_attribution(
        paper,
        ReductionOptions(
            corpus=corpus,
            corpus_license="CC0",
            apply=True,
            output=output,
            candidate_count=1,
            rewrite_provider=provider,
            semantic_backend=_PreservedSemantic(),
            entailment_backend=SelectiveEntailment(),
            audit_runner=_audit,
        ),
    )
    assert output.read_bytes() == paper.read_bytes()
    assert result.metrics.accepted_candidates == 0
    assert any(change.unsupported_new_claims for change in result.changes)


def test_ambiguous_verified_grounding_fails_closed(tmp_path: Path) -> None:
    paper, corpus = _write_fixture(tmp_path)
    preview = improve_attribution(
        paper, ReductionOptions(corpus=corpus, corpus_license="CC0")
    )
    audit = _audit(paper, AuditOptions())
    audit.claim_assessments[0].sources.append(audit.claim_assessments[0].sources[0])
    grounded, manual = build_grounded_reduction_requests(
        audit, preview.risks, preview.plans
    )
    assert grounded == {}
    assert "ambiguous" in next(iter(manual.values()))


def _similarity(percent: float, high: int = 0, medium: int = 0):
    return SimilarityEngineResult([], percent, high, medium, 0, 0, 0, 0)


def test_reduction_metric_percentage_point_and_percent_semantics() -> None:
    metrics = compute_reduction_metrics(_similarity(26.8), _similarity(12.4))
    assert metrics.absolute_reduction_points == pytest.approx(14.4)
    assert metrics.relative_reduction_pct == pytest.approx(53.7313432836)
    unchanged = compute_reduction_metrics(_similarity(10), _similarity(10))
    assert unchanged.relative_reduction_pct == 0
    assert compute_reduction_metrics(_similarity(0), _similarity(0)).relative_reduction_pct is None


def test_docx_single_run_preserves_formatting_and_source(tmp_path: Path) -> None:
    source = tmp_path / "paper.docx"
    output = tmp_path / "revised.docx"
    document = Document()
    paragraph = document.add_paragraph()
    run = paragraph.add_run(ORIGINAL)
    run.bold = True
    document.add_paragraph("References")
    document.add_paragraph("Smith, J. (2024). Trial.")
    document.save(source)
    original_bytes = source.read_bytes()

    records = write_revised_docx(
        source, output, [TextReplacement(ORIGINAL, REWRITE, "p0s0")]
    )

    assert source.read_bytes() == original_bytes
    reopened = Document(output)
    assert reopened.paragraphs[0].text == REWRITE
    assert reopened.paragraphs[0].runs[0].bold is True
    assert reopened.paragraphs[2].text == "Smith, J. (2024). Trial."
    assert records[0].formatting_preserved is True


def test_docx_writer_refuses_existing_output(tmp_path: Path) -> None:
    source = tmp_path / "paper.docx"
    output = tmp_path / "existing.docx"
    document = Document()
    document.add_paragraph(ORIGINAL)
    document.save(source)
    document.save(output)

    with pytest.raises(ValueError, match="already exists"):
        write_revised_docx(
            source, output, [TextReplacement(ORIGINAL, REWRITE, "p0s0")]
        )


def test_docx_cross_run_target_fails_closed(tmp_path: Path) -> None:
    source = tmp_path / "complex.docx"
    output = tmp_path / "complex.revised.docx"
    document = Document()
    paragraph = document.add_paragraph()
    paragraph.add_run("The intervention significantly ").bold = True
    paragraph.add_run("reduced mortality by 12% (Smith, 2024).")
    document.save(source)
    with pytest.raises(UnsafeDocxStructureError, match="crosses runs"):
        write_revised_docx(
            source, output, [TextReplacement(ORIGINAL, REWRITE, "p0s0")]
        )
    assert not output.exists()


def test_end_to_end_docx_apply_rescans_written_output(tmp_path: Path) -> None:
    source = tmp_path / "paper.docx"
    output = tmp_path / "paper.revised.docx"
    corpus = tmp_path / "corpus.txt"
    document = Document()
    document.add_paragraph(ORIGINAL).runs[0].italic = True
    document.add_paragraph("References")
    document.add_paragraph("Smith, J. (2024). Trial.")
    document.save(source)
    corpus.write_text(SOURCE, encoding="utf-8")
    original_bytes = source.read_bytes()

    result = improve_attribution(
        source,
        ReductionOptions(
            corpus=corpus,
            corpus_license="CC0",
            apply=True,
            output=output,
            candidate_count=1,
            rewrite_provider=_RewriteProvider(),
            semantic_backend=_PreservedSemantic(),
            entailment_backend=_PreservedEntailment(),
            audit_runner=_audit,
        ),
    )

    assert source.read_bytes() == original_bytes
    assert Document(output).paragraphs[0].text == REWRITE
    assert Document(output).paragraphs[0].runs[0].italic is True
    assert result.after.overall_similarity_pct < result.before.overall_similarity_pct
    applied = [change for change in result.changes if change.application_status == "applied"]
    assert applied[0].affected_runs == (0,)
    assert applied[0].formatting_preserved is True
