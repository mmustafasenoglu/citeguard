from pathlib import Path

import pytest
from docx import Document

from citeguard.extractor import extract_citations, parse_document, read_paragraphs
from citeguard.retrieval import normalize_doi


def test_extracts_author_year_numbered_and_doi() -> None:
    paragraphs = [
        "Transformers were introduced in 2017 (Vaswani et al., 2017).",
        "Several studies report this effect [12, 13].",
        "See 10.1000/XYZ.123 for the record.",
    ]
    citations = extract_citations(paragraphs)

    assert any(c.authors == "Vaswani et al." and c.year == 2017 for c in citations)
    assert [c.numbered_ref for c in citations if c.numbered_ref] == [12, 13]
    assert any(c.doi == "10.1000/xyz.123" for c in citations)


def test_normalize_doi() -> None:
    assert normalize_doi("https://doi.org/10.1000/ABC.1") == "10.1000/abc.1"
    assert normalize_doi("doi:10.1000/ABC.1") == "10.1000/abc.1"


def test_reads_markdown_fixture() -> None:
    fixture = Path(__file__).parent / "fixtures" / "simple.md"
    paragraphs = read_paragraphs(fixture)
    assert "# Example" in paragraphs
    assert any("Vaswani" in p for p in paragraphs)


def test_rejects_unsupported_file() -> None:
    with pytest.raises(ValueError, match="Unsupported file type"):
        read_paragraphs("paper.pdf")


def test_reads_docx_body_and_table_cells_in_document_order(tmp_path) -> None:
    path = tmp_path / "form.docx"
    document = Document()
    document.add_paragraph("Before table")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "First cell (Smith, 2020)."
    table.cell(0, 1).text = "Second cell"
    merged = table.cell(1, 0).merge(table.cell(1, 1))
    merged.text = "Merged cell"
    document.add_paragraph("After table")
    document.save(path)

    assert read_paragraphs(path) == [
        "Before table",
        "First cell (Smith, 2020).",
        "Second cell",
        "Merged cell",
        "After table",
    ]


def test_accepts_abbreviated_multi_author_markers() -> None:
    citations = extract_citations(["A prior method (Ronneberger vd., 2015) is relevant."])
    assert len(citations) == 1
    assert citations[0].authors == "Ronneberger vd."
    assert citations[0].year == 2015


def test_extracts_narrative_author_year_citations() -> None:
    paragraph = "Filipponi (2018) proposed an index; Ronneberger vd. (2015) introduced U-Net."
    citations = extract_citations([paragraph])

    assert [(citation.authors, citation.year) for citation in citations] == [
        ("Filipponi", 2018),
        ("Ronneberger vd.", 2015),
    ]
    assert [citation.raw_text for citation in citations] == [
        "Filipponi (2018)",
        "Ronneberger vd. (2015)",
    ]
    assert [citation.char_offset for citation in citations] == [0, 36]


def test_extracts_multiple_parenthetical_citations() -> None:
    raw = "(Smith, 2020; Jones vd., 2021)"
    citations = extract_citations([f"Prior work {raw} supports this."])

    assert [(citation.authors, citation.year) for citation in citations] == [
        ("Smith", 2020),
        ("Jones vd.", 2021),
    ]
    assert all(citation.raw_text == raw for citation in citations)
    assert all(citation.char_offset == 11 for citation in citations)


def test_extracts_corporate_author_and_page_locator() -> None:
    citations = extract_citations(
        ["Strategy (T.C. Sanayi ve Teknoloji Bakanlığı, 2025, s. 14) sets priorities."]
    )

    assert len(citations) == 1
    assert citations[0].authors == "T.C. Sanayi ve Teknoloji Bakanlığı"
    assert citations[0].year == 2025


def test_extracts_turkish_and_english_no_date_citations() -> None:
    citations = extract_citations(
        ["Catalogs are available (Google Earth Engine Data Catalog, t.y.; Copernicus, n.d.)."]
    )

    assert [citation.authors for citation in citations] == [
        "Google Earth Engine Data Catalog",
        "Copernicus",
    ]
    assert all(citation.year is None and citation.no_date for citation in citations)


def test_does_not_treat_ordinary_parentheses_as_citations() -> None:
    citations = extract_citations(["The selected band (B8A, 20 meters) is resampled."])

    assert citations == []


def test_parse_document_does_not_count_bibliography_dois_as_citations(tmp_path) -> None:
    path = tmp_path / "paper.txt"
    path.write_text(
        "A supported claim (Smith, 2020).\n\n"
        "References\n\n"
        "Smith, J. (2020). Paper. 10.1000/example",
        encoding="utf-8",
    )
    parsed = parse_document(path)
    assert len(parsed.citations) == 1
    assert parsed.citations[0].authors == "Smith"
